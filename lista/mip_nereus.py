# -*- coding: utf-8 -*-
"""
mip_nereus.py — carga das matrizes nacionais "MIP-BR <ano> (Nível 68).xlsm"
(NEREUS/CECEG, layout espelhado na publicação da MIP do IBGE) para a Lista de
Exercícios da disciplina (Prof. Celso Bissoli Sessa, 2026/1).

Convenções herdadas do pipeline do repositório (pesquisa/18_trajetoria_n68.py):
  - aba '13' = matriz A de coeficientes técnicos dos insumos NACIONAIS (68x68);
  - aba '11' = fluxos atividade x atividade da produção nacional a preço básico
    (Z nacional, demanda final por componente e bloco de VA/ocupações);
  - aba '12' = idem para produtos IMPORTADOS (uso intermediário importado e
    demanda final importada, base dos coeficientes de importação).

Como os arquivos reais variam em offsets, a carga é dirigida por RÓTULOS
(autodetecção do bloco de setores e das linhas/colunas nomeadas), com
`--inspect` para depurar o layout de qualquer arquivo antes de resolver.
Valores em R$ milhões correntes do ano da matriz.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field

import numpy as np
import openpyxl


# --------------------------------------------------------------------------- #
# utilidades
# --------------------------------------------------------------------------- #
def norm(s) -> str:
    """minúsculas, sem acento, espaços colapsados — para casar rótulos."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def num(v) -> float:
    return float(v) if isinstance(v, (int, float)) else 0.0


def grade(ws, max_row=None, max_col=None):
    """matriz de valores da planilha (lista de listas, indexada de 0)."""
    return [list(r) for r in ws.iter_rows(min_row=1, max_row=max_row,
                                          max_col=max_col, values_only=True)]


# --------------------------------------------------------------------------- #
# contêiner de um ano da MIP
# --------------------------------------------------------------------------- #
COMPONENTES_DF = ["exportacao", "governo", "isflsf", "familias", "fbcf", "estoque"]

# padrões (via norm) que identificam cada componente de demanda final no cabeçalho
PADROES_DF = {
    "exportacao": ("exportacao",),
    "governo":    ("governo", "administracao publica"),
    "isflsf":     ("isflsf", "instituicoes sem fins"),
    "familias":   ("familias",),
    "fbcf":       ("formacao bruta", "fbcf"),
    "estoque":    ("variacao de estoque", "estoque"),
}

# padrões das linhas do bloco de valor adicionado / totais abaixo de Z
PADROES_LINHAS = {
    "importacao":  ("importacao",),
    "impostos":    ("impostos",),
    "va":          ("valor adicionado",),
    "rem":         ("remuneracoes",),
    "salarios":    ("salarios",),
    "vp":          ("valor da producao", "valor de producao", "valor bruto da producao"),
    "ocup":        ("fator trabalho", "ocupacoes", "pessoal ocupado", "emprego"),
}


@dataclass
class MIPAno:
    ano: int
    nomes: list[str]
    cod: list[str]
    Z: np.ndarray                      # consumo intermediário NACIONAL (n x n)
    y: dict[str, np.ndarray]           # demanda final nacional por componente (n)
    x: np.ndarray                      # valor da produção (n)
    rem: np.ndarray                    # remunerações (n)
    va: np.ndarray                     # valor adicionado bruto (n)
    ocup: np.ndarray                   # fator trabalho / ocupações (n)
    imp_int: np.ndarray                # importação de uso intermediário por setor comprador (n)
    y_imp: dict[str, np.ndarray] = field(default_factory=dict)  # demanda final IMPORTADA por componente (n)
    proveniencia: dict = field(default_factory=dict)            # aba/faixa de cada peça (p/ Metodologia)

    @property
    def n(self) -> int:
        return len(self.nomes)

    def y_total(self) -> np.ndarray:
        return np.sum([self.y[c] for c in COMPONENTES_DF], axis=0)

    def validar(self, estrito: bool = True) -> list[str]:
        """checagens contábeis; devolve avisos (e levanta erro se estrito e algo grave)."""
        avisos = []
        n = self.n
        x_row = self.Z.sum(1) + self.y_total()
        err = np.max(np.abs(self.x - x_row) / np.where(self.x == 0, 1, np.abs(self.x)))
        if err > 0.02:
            avisos.append(f"balanço de linha x = Z·1 + y com desvio máximo de {err:.1%}")
        A = self.Z / np.where(self.x == 0, 1, self.x)[None, :]
        if A.sum(0).max() >= 1:
            avisos.append(f"soma de coluna de A chega a {A.sum(0).max():.3f} (>=1)")
        for nome_v, v in [("rem", self.rem), ("va", self.va), ("ocup", self.ocup)]:
            if v.shape != (n,):
                avisos.append(f"vetor {nome_v} com dimensão {v.shape} != ({n},)")
        if estrito and any("dimensão" in a for a in avisos):
            raise ValueError("; ".join(avisos))
        return avisos


# --------------------------------------------------------------------------- #
# autodetecção do bloco de setores
# --------------------------------------------------------------------------- #
def achar_bloco_setores(G, min_run=5):
    """
    Procura a coluna de NOMES de setor: sequência contígua de células de texto
    (>= min_run) que não sejam rótulos de total. Entre candidatas de comprimento
    parecido (códigos e nomes correm lado a lado), prefere a de texto mais LONGO
    (nomes), não a de códigos. Devolve (linha_ini, linha_fim, col_nomes), base 0.
    """
    candidatas = []
    nlin = len(G)
    ncol = max(len(r) for r in G) if G else 0
    for c in range(min(ncol, 8)):                      # nomes ficam nas primeiras colunas
        r = 0
        while r < nlin:
            v = G[r][c] if c < len(G[r]) else None
            if isinstance(v, str) and norm(v) and not _eh_total(v):
                r0 = r
                textos = []
                while r < nlin:
                    v = G[r][c] if c < len(G[r]) else None
                    if isinstance(v, str) and norm(v) and not _eh_total(v):
                        textos.append(str(v)); r += 1
                    else:
                        break
                if len(textos) >= min_run:
                    media = sum(len(t) for t in textos) / len(textos)
                    candidatas.append((r0, r - 1, c, len(textos), media))
            else:
                r += 1
    if not candidatas:
        raise ValueError("não achei bloco contíguo de nomes de setor — use --inspect")
    run_max = max(c[3] for c in candidatas)
    finalistas = [c for c in candidatas if c[3] >= 0.9 * run_max]
    # setores vêm ACIMA do bloco de VA/totais: menor linha inicial primeiro;
    # empate (códigos × nomes lado a lado) decidido pelo texto médio mais longo
    r0, r1, c, _, _ = min(finalistas, key=lambda t: (t[0], -t[4]))
    return r0, r1, c


def _eh_total(v) -> bool:
    s = norm(v)
    return s.startswith("total") or s in ("demanda final", "demanda total",
                                          "consumo intermediario", "usos")


def _bloco_numerico(G, r0, r1, col_ini, n):
    """primeiras n colunas 'densamente numéricas' a partir de col_ini nas linhas r0..r1
    (pulando colunas de índice sequencial 1..n, comuns nos arquivos do IBGE/NEREUS)."""
    ncol = max(len(r) for r in G)
    cols = []
    c = col_ini
    while c < ncol and len(cols) < n:
        vals = [G[r][c] if c < len(G[r]) else None for r in range(r0, r1 + 1)]
        dens = sum(isinstance(v, (int, float)) for v in vals) / (r1 - r0 + 1)
        eh_rampa = all(isinstance(v, (int, float)) and abs(v - (k + 1)) < 1e-9
                       for k, v in enumerate(vals))
        if dens >= 0.5 and not eh_rampa:
            cols.append(c)
        elif cols:                                     # bloco contíguo interrompido
            break
        c += 1
    if len(cols) < n:
        raise ValueError(f"esperava {n} colunas numéricas a partir da coluna "
                         f"{col_ini + 1}; achei {len(cols)} — use --inspect")
    return cols


def _cabecalho(G, r0, cols_alvo, alcance=6):
    """texto de cabeçalho de cada coluna: concatena até `alcance` linhas acima do bloco."""
    out = {}
    for c in cols_alvo:
        partes = []
        for r in range(max(0, r0 - alcance), r0):
            v = G[r][c] if c < len(G[r]) else None
            if isinstance(v, str) and v.strip():
                partes.append(v)
        out[c] = norm(" ".join(partes))
    return out


# --------------------------------------------------------------------------- #
# carga de uma aba de fluxos (nacional '11' ou importados '12')
# --------------------------------------------------------------------------- #
def _carregar_fluxos(ws, rotulo_aba):
    G = grade(ws)
    r0, r1, c_nome = achar_bloco_setores(G)
    n = r1 - r0 + 1
    nomes = [str(G[r][c_nome]).replace("\n", " ").strip() for r in range(r0, r1 + 1)]
    cods = []
    for r in range(r0, r1 + 1):                        # código: 1ª célula não vazia à esquerda
        cod = ""
        for c in range(c_nome):
            v = G[r][c] if c < len(G[r]) else None
            if v is not None and str(v).strip():
                cod = str(v).strip()
        cods.append(cod)

    cols_Z = _bloco_numerico(G, r0, r1, c_nome + 1, n)
    Z = np.array([[num(G[r][c] if c < len(G[r]) else None) for c in cols_Z]
                  for r in range(r0, r1 + 1)])

    # componentes de demanda final: colunas à direita de Z, casadas pelo cabeçalho
    ncol = max(len(r) for r in G)
    cols_dir = list(range(cols_Z[-1] + 1, ncol))
    cab = _cabecalho(G, r0, cols_dir)
    ycomp = {}
    for comp, padroes in PADROES_DF.items():
        for c in cols_dir:
            if any(p in cab[c] for p in padroes) and comp not in ycomp:
                ycomp[comp] = np.array([num(G[r][c] if c < len(G[r]) else None)
                                        for r in range(r0, r1 + 1)])
                break

    # linhas rotuladas abaixo do bloco (VA, remunerações, VP, ocupações, importação…)
    linhas = {}
    for r in range(r1 + 1, len(G)):
        rot = ""
        for c in range(0, c_nome + 1):
            v = G[r][c] if c < len(G[r]) else None
            if isinstance(v, str) and v.strip():
                rot = norm(v)
        if not rot:
            continue
        for chave, padroes in PADROES_LINHAS.items():
            if any(rot.startswith(p) or p in rot for p in padroes) and chave not in linhas:
                linhas[chave] = np.array([num(G[r][c] if c < len(G[r]) else None)
                                          for c in cols_Z])
    prov = {
        "aba": rotulo_aba,
        "linhas_setores": (r0 + 1, r1 + 1),
        "col_nomes": c_nome + 1,
        "cols_Z": (cols_Z[0] + 1, cols_Z[-1] + 1),
        "df_encontrada": sorted(ycomp),
        "linhas_encontradas": sorted(linhas),
    }
    return nomes, cods, Z, ycomp, linhas, prov


def _achar_aba(wb, candidatos):
    for nome in candidatos:
        if nome in wb.sheetnames:
            return nome
    for nome in wb.sheetnames:                          # tolera '11 ' / 'Tab 11' etc.
        base = norm(nome).replace("tab", "").strip()
        if base in [norm(c) for c in candidatos]:
            return nome
    return None


def carregar_mip(caminho: str, ano: int) -> MIPAno:
    """Carrega um 'MIP-BR <ano> (Nível 68).xlsm/xlsx' e devolve MIPAno validado."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        aba_nac = _achar_aba(wb, ["11", "03"])
        if aba_nac is None:
            raise ValueError(f"aba de fluxos nacionais ('11') não encontrada em "
                             f"{wb.sheetnames} — use --inspect")
        nomes, cods, Z, ycomp, linhas, prov = _carregar_fluxos(wb[aba_nac], aba_nac)
        n = len(nomes)

        faltam = [c for c in COMPONENTES_DF if c not in ycomp]
        if faltam:
            raise ValueError(f"componentes de demanda final ausentes na aba {aba_nac}: "
                             f"{faltam} (achados: {sorted(ycomp)}) — use --inspect")

        y_imp, imp_int, prov_imp = {}, np.zeros(n), None
        aba_imp = _achar_aba(wb, ["12", "04"])
        if aba_imp is not None:
            _, _, Zm, ycomp_m, _, prov_imp = _carregar_fluxos(wb[aba_imp], aba_imp)
            if Zm.shape == (n, n):
                imp_int = Zm.sum(0)                    # importado intermediário por comprador
                y_imp = ycomp_m
        if not np.any(imp_int) and "importacao" in linhas:
            imp_int = linhas["importacao"]             # fallback: linha 'Importação'

        x = linhas.get("vp", np.zeros(n))
        if not np.any(x):
            x = Z.sum(1) + np.sum([ycomp[c] for c in COMPONENTES_DF], axis=0)
        rem = linhas.get("rem", linhas.get("salarios", np.zeros(n)))
        va = linhas.get("va", np.zeros(n))
        ocup = linhas.get("ocup", np.zeros(n))

        mip = MIPAno(ano=ano, nomes=nomes, cod=cods, Z=Z, y=ycomp, x=x,
                     rem=rem, va=va, ocup=ocup, imp_int=imp_int, y_imp=y_imp,
                     proveniencia={"arquivo": caminho, "nacional": prov,
                                   "importados": prov_imp})
        for chave, v in [("remunerações", rem), ("valor adicionado", va),
                         ("ocupações", ocup)]:
            if not np.any(v):
                raise ValueError(f"linha de {chave} não encontrada na aba {aba_nac} "
                                 f"(linhas achadas: {prov['linhas_encontradas']}) — use --inspect")
        return mip
    finally:
        wb.close()


# --------------------------------------------------------------------------- #
# inspeção de layout
# --------------------------------------------------------------------------- #
def inspecionar(caminho: str) -> str:
    """Relatório do layout de cada aba: dimensões, rótulos de linha e cabeçalhos."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    saida = [f"arquivo: {caminho}", f"abas: {wb.sheetnames}"]
    for nome in wb.sheetnames:
        ws = wb[nome]
        G = grade(ws, max_row=min(ws.max_row or 0, 200), max_col=min(ws.max_column or 0, 120))
        saida.append(f"\n--- aba '{nome}' ({ws.max_row}x{ws.max_column}) ---")
        try:
            r0, r1, c = achar_bloco_setores(G)
            saida.append(f"  bloco de setores: linhas {r0+1}-{r1+1} ({r1-r0+1} setores), "
                         f"nomes na coluna {c+1}")
            saida.append(f"  1º setor: {G[r0][c]!r}   último: {G[r1][c]!r}")
            for r in range(r1 + 1, min(r1 + 25, len(G))):
                rot = next((str(G[r][cc]) for cc in range(0, c + 1)
                            if cc < len(G[r]) and isinstance(G[r][cc], str) and G[r][cc].strip()), "")
                if rot:
                    saida.append(f"  linha {r+1}: {rot[:80]!r}")
        except ValueError as e:
            saida.append(f"  (sem bloco de setores: {e})")
    wb.close()
    return "\n".join(saida)


# --------------------------------------------------------------------------- #
# gerador sintético (selftest): mesmo layout '11'/'12'/'13' em miniatura
# --------------------------------------------------------------------------- #
def escrever_sintetico(caminho: str, ano: int, semente: int = 7, n: int = 6) -> None:
    """Escreve um arquivo no layout NEREUS com uma economia sintética balanceada."""
    rng = np.random.default_rng(semente + ano)
    nomes = [f"Setor sintético {i+1:02d}" for i in range(n)]
    x = rng.uniform(1e5, 5e5, n)
    A = rng.uniform(0.01, 0.10, (n, n)); np.fill_diagonal(A, rng.uniform(0.04, 0.10, n))
    Z = A * x[None, :]
    while (x - Z.sum(1)).min() <= 0.05 * x.min():       # garante demanda final positiva
        A *= 0.8
        Z = A * x[None, :]
    m_coef = rng.uniform(0.02, 0.08, n)
    Zm = (m_coef * x)[None, :] * rng.dirichlet(np.ones(n), n).T
    va = x - Z.sum(0) - Zm.sum(0)
    rem = va * rng.uniform(0.35, 0.6, n)
    ocup = x * rng.uniform(5, 40, n) / 1e3
    resto = x - Z.sum(1)                                # demanda final nacional total
    pesos = rng.dirichlet(np.ones(6), n)                # 6 componentes
    Ynac = pesos * resto[:, None]
    Yimp = rng.uniform(0, 0.08, (n, 6)) * resto[:, None] / 6

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cabec = ["Exportação de bens e serviços", "Consumo do governo",
             "Consumo das ISFLSF", "Consumo das famílias",
             "Formação bruta de capital fixo", "Variação de estoque"]

    def aba_fluxos(tag, M, Y, com_va):
        ws = wb.create_sheet(tag)
        ws.cell(2, 2, f"Tabela sintética {tag} — ano {ano} (R$ milhões)")
        for k, h in enumerate(cabec):
            ws.cell(4, 3 + n + k, h)
        for i in range(n):
            ws.cell(5 + i, 1, f"{i+1:04d}")
            ws.cell(5 + i, 2, nomes[i])
            for j in range(n):
                ws.cell(5 + i, 3 + j, float(M[i, j]))
            for k in range(6):
                ws.cell(5 + i, 3 + n + k, float(Y[i, k]))
        if com_va:
            blocos = [("Importação", Zm.sum(0)), ("Impostos líquidos", 0.0 * x),
                      ("Valor adicionado bruto ( PIB )", va), ("Remunerações", rem),
                      ("Valor da produção", x), ("Fator trabalho (ocupações)", ocup)]
            for b, (rot, vet) in enumerate(blocos):
                ws.cell(6 + n + b, 2, rot)
                for j in range(n):
                    ws.cell(6 + n + b, 3 + j, float(np.asarray(vet).reshape(-1)[j] if np.ndim(vet) else vet))

    aba_fluxos("11", Z, Ynac, com_va=True)
    aba_fluxos("12", Zm, Yimp, com_va=False)
    ws13 = wb.create_sheet("13")
    ws13.cell(2, 2, "Matriz de coeficientes técnicos dos insumos nacionais (sintética)")
    for i in range(n):
        ws13.cell(7 + i, 2, f"{i+1:04d}")
        ws13.cell(7 + i, 3, nomes[i])
        for j in range(n):
            ws13.cell(7 + i, 5 + j, float(A[i, j]))
    wb.save(caminho)
