# -*- coding: utf-8 -*-
"""
auditoria.py — auditoria HARD da resolução da lista contra o MATERIAL DO PROFESSOR.

O gabarito é a própria pasta MIP-BR (Nível 68): além das TRUs/MIP, os arquivos
trazem as abas de cálculo do Prof. Celso (inversa de Leontief, Ghosh, índices de
ligação, multiplicadores tipo I/II, modelo fechado, geradores). Este script
extrai essas abas e confronta, célula a célula, com o que o solver calcula —
falhando (exit 1) se qualquer confronto exceder a tolerância.

Abas auditadas (nos arquivos de 2010 e 2020):
  13     (I−A)⁻¹ ...................... inversa de Leontief B
  Ghosh  (I−Â)⁻¹ ...................... inversa de Ghosh G
  14     índices de ligação (aberto) .. Rasmussen-Hirschman para trás/frente
  15     geradores/multiplicadores .... impacto tipo I (ocup., rem., VAB, import.)
  22     geradores (modelo fechado) ... impacto tipo II
  23     multiplicadores tipo I e II .. produção aberto/fechado

Uso:  python lista/auditoria.py       (roda sobre dados/MIP-BR {2010,2020}...)
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from mip_nereus import carregar_mip                      # noqa: E402
from resolver_lista import Sistema, q6_extracao          # noqa: E402

ARQ = {2010: "dados/MIP-BR 2010 (Nível 68).xlsm",
       2020: "dados/MIP-BR 2020 (Nível 68).xlsm"}
TOL = 1e-9                       # tolerância dos confrontos (identidade numérica)
N = 68


def _bloco(ws, r0, nlin, c0=5, ncol=N):
    return np.array([[v if isinstance(v, (int, float)) else 0.0
                      for v in row[c0 - 1:c0 - 1 + ncol]]
                     for row in ws.iter_rows(min_row=r0, max_row=r0 + nlin - 1,
                                             max_col=c0 - 1 + ncol, values_only=True)])


def _rot(ws, rot, r0, r1):
    for r in range(r0, r1):
        v = ws.cell(r, 2).value
        if v and rot.lower() in str(v).strip().lower():
            return r
    return None


def auditar(ano: int, resultados: list):
    m = carregar_mip(ARQ[ano], ano)
    s = Sistema(m)
    wb = openpyxl.load_workbook(ARQ[ano], read_only=False, data_only=True)

    def reg(nome, obtido_vs_prof_maxdiff, ok=None):
        ok = (obtido_vs_prof_maxdiff <= TOL) if ok is None else ok
        resultados.append((ano, nome, obtido_vs_prof_maxdiff, ok))

    # 1) inversa de Leontief B  (aba 13, bloco 'Matriz (I-A)-1' em r=226)
    B13 = _bloco(wb["13"], 226, N)
    reg("B = (I−A)⁻¹  vs aba 13", float(np.abs(s.B - B13).max()))

    # 2) inversa de Ghosh G  (aba Ghosh, bloco '(I-A)-1' em r=223)
    G = _bloco(wb["Ghosh"], 223, N)
    reg("G = (I−Â)⁻¹  vs aba Ghosh", float(np.abs(s.G - G).max()))

    # 3) ligações Rasmussen-Hirschman  (aba 14: E=trás bruto, F=frente bruto,
    #    G=trás norm, H=frente norm)
    lig = np.array([[v if isinstance(v, (int, float)) else np.nan for v in r[4:8]]
                    for r in wb["14"].iter_rows(min_row=7, max_row=6 + N, max_col=8,
                                                values_only=True)])
    bl, fl = s.rasmussen()
    reg("ligação p/ trás (norm.)  vs aba 14", float(np.nanmax(np.abs(lig[:, 2] - bl))))
    reg("ligação p/ frente (norm.) vs aba 14", float(np.nanmax(np.abs(lig[:, 3] - fl))))

    # 4) geradores TIPO I (aba 15, bloco 'Geradores' em r=26..39; dados na coluna E=5)
    ws15 = wb["15"]
    ger = {rot: np.array([ws15.cell(_rot(ws15, rot, 26, 40), 5 + j).value for j in range(N)])
           for rot in ("Fator Trabalho", "Remunerações", "Valor Adicionado", "Importação")}
    emp_I = (s.w_emp[:, None] * s.B).sum(0)
    ren_I = (s.v_renda[:, None] * s.B).sum(0)
    va_I = (s.v_va[:, None] * s.B).sum(0)
    imp_I = (s.m_coef[:, None] * s.B).sum(0)
    reg("gerador emprego  (tipo I) vs aba 15", float(np.abs(ger["Fator Trabalho"] - emp_I).max()))
    reg("gerador renda    (tipo I) vs aba 15", float(np.abs(ger["Remunerações"] - ren_I).max()))
    reg("gerador VAB      (tipo I) vs aba 15", float(np.abs(ger["Valor Adicionado"] - va_I).max()))
    reg("gerador importação(tipo I) vs aba 15", float(np.abs(ger["Importação"] - imp_I).max()))

    # 5) geradores TIPO II (aba 22, modelo fechado; dados na coluna E=5)
    ws22 = wb["22"]
    ger2 = {rot: np.array([ws22.cell(_rot(ws22, rot, 23, 40), 5 + j).value for j in range(N)])
            for rot in ("Fator Trabalho", "Remunerações", "Valor Adicionado")}
    B2n = s.B2[:N, :N]
    emp_II = (s.w_emp[:, None] * B2n).sum(0)
    ren_II = (s.v_renda[:, None] * B2n).sum(0)
    va_II = (s.v_va[:, None] * B2n).sum(0)
    reg("gerador emprego  (tipo II) vs aba 22", float(np.abs(ger2["Fator Trabalho"] - emp_II).max()))
    reg("gerador renda    (tipo II) vs aba 22", float(np.abs(ger2["Remunerações"] - ren_II).max()))
    reg("gerador VAB      (tipo II) vs aba 22", float(np.abs(ger2["Valor Adicionado"] - va_II).max()))

    # 6) multiplicadores de produção tipo I e II (aba 23: E=I, F=II)
    t23 = np.array([[v if isinstance(v, (int, float)) else np.nan for v in r[4:7]]
                    for r in wb["23"].iter_rows(min_row=7, max_row=6 + N, max_col=7,
                                                values_only=True)])
    reg("mult. produção tipo I  vs aba 23", float(np.nanmax(np.abs(t23[:, 0] - s.B.sum(0)))))
    reg("mult. produção tipo II vs aba 23", float(np.nanmax(np.abs(t23[:, 1] - B2n.sum(0)))))

    wb.close()
    return m, s


def main():
    faltando = [p for p in ARQ.values() if not (RAIZ / p).exists()]
    if faltando:
        sys.exit(f"[ERRO] arquivos ausentes: {faltando}\nColoque as MIPs em dados/ (ver lista/README.md).")

    resultados: list = []
    for ano in (2010, 2020):
        auditar(ano, resultados)

    print("=" * 78)
    print("AUDITORIA DA LISTA vs MATERIAL DO PROFESSOR (abas de cálculo da MIP-BR)")
    print(f"tolerância: {TOL:.0e}  ·  {N} setores")
    print("=" * 78)
    falhas = 0
    for ano, nome, diff, ok in resultados:
        status = "OK  " if ok else "FALHA"
        if not ok:
            falhas += 1
        print(f"[{status}] {ano} · {nome:42s} maxdiff = {diff:.2e}")
    print("-" * 78)
    if falhas:
        print(f"[FALHA] {falhas} confronto(s) acima da tolerância.")
        sys.exit(1)
    print(f"[OK] {len(resultados)} confrontos reproduzem o material do professor "
          f"(desvio máximo < {TOL:.0e}).")


if __name__ == "__main__":
    main()
