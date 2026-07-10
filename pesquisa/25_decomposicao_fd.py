# -*- coding: utf-8 -*-
"""
25_decomposicao_fd.py — C4: QUEM PUXA a economia capixaba? Decomposicao da producao
e do emprego do ES pelos componentes da demanda final, x^(k) = B y^(k) (aditiva exata).

Componentes (Tabela `tab:decomp` do artigo):
  - Demanda final do RESTANTE DO BRASIL (bloco RB, exceto exportacoes internacionais)
  - Exportacoes internacionais (colunas Export dos DOIS blocos, ES e RB)
  - Consumo das familias ES · Governo+ISFLSF ES · FBCF ES · Variacao de estoques ES

Fonte: MIP-ES-BR (2008).xlsx, aba 'ES-BR'. Layout (de 01_es_br_base.py):
  Z 52x52 = linhas 7-58 x cols 5-56; FD cols 58-69
  (58-63 = dem. final ES: Export, Gov, ISFLSF, Familias, FBCF, VarEst; 64-69 = RB, idem)
  VBP linha 80; Ocupacoes linha 81.

NOTA DE PROVENIENCIA: reconstrucao fiel do script local original (nao versionado a
tempo), seguindo as convencoes de 01/17. Um bloco de AUTOVERIFICACAO no final compara
cada numero com a Tabela publicada (34,2/30,4/19,7/13,7/7,6/-1,0 bi; 61,7% producao
externa vs 47,5% emprego) e FALHA se divergir. Se a sua versao original existir,
prefira-a e compare as saidas.
"""
import os, csv, sys
import numpy as np
import openpyxl
if hasattr(sys.stdout, "reconfigure"): sys.stdout.reconfigure(encoding="utf-8")

DATA = sys.argv[1] if len(sys.argv) > 1 else \
       r"C:/Users/DELL/Downloads/Material IO/Matrizes/MIP-ES-BR (2008).xlsx"
OUT  = os.path.join(os.path.dirname(os.path.abspath(__file__)), "outputs")
os.makedirs(OUT, exist_ok=True)
if not os.path.exists(DATA):
    sys.exit(f"[ERRO] matriz nao encontrada: {DATA}\n"
             "Passe o caminho como argumento: python 25_decomposicao_fd.py <caminho do xlsx>")
num = lambda v: float(v) if isinstance(v, (int, float)) else 0.0

wb = openpyxl.load_workbook(DATA, data_only=True); ws = wb["ES-BR"]
names = [str(ws.cell(r, 2).value).strip() for r in range(7, 33)]
Z   = np.array([[num(ws.cell(r, c).value) for c in range(5, 57)] for r in range(7, 59)])
FD  = np.array([[num(ws.cell(r, c).value) for c in range(58, 70)] for r in range(7, 59)])
vbp = np.array([num(ws.cell(80, c).value) for c in range(5, 57)])
emp = np.array([num(ws.cell(81, c).value) for c in range(5, 57)])
wb.close()

n = 52; L = np.arange(0, 26)
x = np.where(vbp == 0, Z.sum(1) + FD.sum(1), vbp)
xs = np.where(x == 0, 1.0, x)
A = Z / xs[None, :]
B = np.linalg.inv(np.eye(n) - A)
w_emp = emp / xs                       # ocupacoes por R$ mi de producao

# --- particao da demanda final em componentes (colunas 0..11 do array FD) ---
# 0..5 = ES (Export, Gov, ISFLSF, Familias, FBCF, VarEst); 6..11 = RB (idem)
COMP = [
    ("Demanda final do restante do Brasil", [7, 8, 9, 10, 11]),   # RB exceto Export
    ("Exportações (resto do mundo)",        [0, 6]),              # Export ES + Export RB
    ("Consumo das famílias ES",             [3]),
    ("Governo + ISFLSF ES",                 [1, 2]),
    ("Formação bruta de capital ES",        [4]),
    ("Variação de estoques ES",             [5]),
]
assert sorted(sum((c for _, c in COMP), [])) == list(range(12)), "particao nao cobre as 12 colunas"

prod_k = []; emp_k = []
for nome, cols in COMP:
    yk = FD[:, cols].sum(1)
    xk = B @ yk
    prod_k.append(float(xk[L].sum()))                 # producao dos setores do ES
    emp_k.append(float((w_emp[L] * xk[L]).sum()))     # ocupacoes nos setores do ES
prod_k = np.array(prod_k); emp_k = np.array(emp_k)
prod_tot, emp_tot = prod_k.sum(), emp_k.sum()

# aditividade exata: soma dos componentes = producao/emprego totais do ES
x_full = B @ FD.sum(1)
assert abs(prod_tot - x_full[L].sum()) < 1e-6 * prod_tot, "aditividade da producao quebrou"
assert abs(emp_tot - (w_emp[L] * x_full[L]).sum()) < 1e-6 * max(emp_tot, 1), "aditividade do emprego quebrou"

# parcela EXTERNA por setor (RB + exportacoes), p/ o contraste setorial do texto
y_ext = FD[:, [7, 8, 9, 10, 11, 0, 6]].sum(1)
x_ext = B @ y_ext
pct_ext_setor = x_ext[L] / np.where(x_full[L] == 0, 1, x_full[L]) * 100

ext_prod = (prod_k[0] + prod_k[1]) / prod_tot * 100
ext_emp  = (emp_k[0] + emp_k[1]) / emp_tot * 100

print("=" * 78)
print("C4 — QUEM PUXA A ECONOMIA DO ES? Decomposicao pela demanda final (2008)")
print("=" * 78)
print(f"{'componente':38s} {'prod R$ bi':>10s} {'%':>6s} {'emp mil':>9s} {'%':>6s}")
for k, (nome, _) in enumerate(COMP):
    print(f"{nome:38s} {prod_k[k]/1e3:10.1f} {prod_k[k]/prod_tot*100:6.1f} "
          f"{emp_k[k]/1e3:9.1f} {emp_k[k]/emp_tot*100:6.1f}")
print(f"{'TOTAL':38s} {prod_tot/1e3:10.1f} {100.0:6.1f} {emp_tot/1e3:9.1f} {100.0:6.1f}")
print("-" * 78)
print(f"demanda EXTERNA ao estado: {ext_prod:.1f}% da producao · {ext_emp:.1f}% do emprego")
o = np.argsort(pct_ext_setor)[::-1]
print("setores mais puxados de FORA:",
      ", ".join(f"{names[j][:20]} ({pct_ext_setor[j]:.1f}%)" for j in o[:4]))
print("=" * 78)

# ------------------- AUTOVERIFICACAO contra a Tabela publicada -------------------
ALVO = {  # (producao R$ bi, % prod, emprego mil, % emp) — Tabela `tab:decomp`
    "Demanda final do restante do Brasil": (34.2, 32.7, 477.8, 29.6),
    "Exportações (resto do mundo)":        (30.4, 29.1, 287.8, 17.9),
    "Consumo das famílias ES":             (19.7, 18.8, 414.5, 25.7),
    "Governo + ISFLSF ES":                 (13.7, 13.1, 259.4, 16.1),
    "Formação bruta de capital ES":        (7.6,  7.3,  163.7, 10.2),
    "Variação de estoques ES":             (-1.0, -1.0, 8.6,   0.5),
}
fails = []
def chk(rot, got, alvo, tol):
    ok = abs(got - alvo) <= tol
    if not ok: fails.append(f"{rot}: obtido {got:.1f} vs publicado {alvo:.1f}")
    return "OK " if ok else "FAIL"
print("AUTOVERIFICACAO (vs Tabela publicada; tolerancia 0,1 bi / 0,2 pp / 1 mil):")
for k, (nome, _) in enumerate(COMP):
    a = ALVO[nome]
    s1 = chk(f"{nome} prod", prod_k[k]/1e3, a[0], 0.1)
    s2 = chk(f"{nome} %prod", prod_k[k]/prod_tot*100, a[1], 0.2)
    s3 = chk(f"{nome} emp", emp_k[k]/1e3, a[2], 1.0)
    s4 = chk(f"{nome} %emp", emp_k[k]/emp_tot*100, a[3], 0.2)
    print(f"  [{s1}{s2}{s3}{s4}] {nome}")
print(f"  [{chk('total prod', prod_tot/1e3, 104.6, 0.2)}] total producao 104,6 bi | "
      f"[{chk('total emp', emp_tot/1e3, 1611.7, 2.0)}] total emprego 1.611,7 mil")
print(f"  [{chk('externa prod', ext_prod, 61.7, 0.2)}] externa 61,7% prod | "
      f"[{chk('externa emp', ext_emp, 47.5, 0.2)}] 47,5% emp")
if fails:
    print("\n[FAIL] a reconstrucao NAO reproduz a tabela publicada nos itens:")
    for f_ in fails: print("   -", f_)
    print("=> confira o mapeamento das colunas de demanda final (COMP) ou use o script original.")
else:
    print("\n[OK] todos os numeros publicados reproduzidos.")

with open(os.path.join(OUT, "decomposicao_fd.csv"), "w", newline="", encoding="utf-8") as f:
    wr = csv.writer(f)
    wr.writerow(["componente", "producao_mi", "producao_pct", "emprego", "emprego_pct"])
    for k, (nome, _) in enumerate(COMP):
        wr.writerow([nome, f"{prod_k[k]:.1f}", f"{prod_k[k]/prod_tot*100:.2f}",
                     f"{emp_k[k]:.1f}", f"{emp_k[k]/emp_tot*100:.2f}"])
    wr.writerow(["TOTAL", f"{prod_tot:.1f}", "100.00", f"{emp_tot:.1f}", "100.00"])
with open(os.path.join(OUT, "decomposicao_fd_setorial.csv"), "w", newline="", encoding="utf-8") as f:
    wr = csv.writer(f)
    wr.writerow(["setor", "pct_producao_externa"])
    for j in range(26):
        wr.writerow([names[j], f"{pct_ext_setor[j]:.2f}"])
print("salvos: decomposicao_fd.csv, decomposicao_fd_setorial.csv")
sys.exit(1 if fails else 0)
